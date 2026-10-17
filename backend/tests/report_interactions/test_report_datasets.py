import io
import json
import zipfile
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select

from app import usage
from app.config import settings
from app.main import app
from app.models import AppUser, ReportDataset, SalesTransaction
from app.services import auth, report_datasets as service
from app.llm import LLMOutputError


CSV = '編號,客戶,日期,金額,備註\n001, 甲店 ,2031/04/01,"1,200",正常\n002,乙店,2031/04/02,300,正常\n003,乙店,2031/04/02,-100,退貨\n004,丙店,2031/04/03,,空白\n'


@pytest.fixture
def dataset_session(interaction_session):
    ReportDataset.__table__.create(interaction_session.connection(), checkfirst=True)
    return interaction_session


def client(session, user="S1"):
    return TestClient(app), {"Authorization": f"Bearer {auth.create_token(session.get(AppUser, user))}"}


def upload(session, content=CSV, filename="銷售.csv", **options):
    api, headers = client(session)
    response = api.post("/api/reports/datasets", headers=headers, files={"file": (filename, content.encode() if isinstance(content, str) else content)}, data=options)
    assert response.status_code == 201, response.text
    return api, headers, response.json()


def chart(**patch):
    return {"id": "test-chart", "title": "金額分析", "visual": "bar", "group_by": "c2", "value": "c4", "aggregation": "sum", **patch}


def test_imports_belong_to_logged_in_account_even_with_same_demo_role(dataset_session, monkeypatch):
    monkeypatch.setattr(settings(), "demo_role_switch", True)
    monkeypatch.setitem(auth.DEMO_ROLE_USERS, "sales", "S1")
    api = TestClient(app)
    headers = {
        account: {"Authorization": f"Bearer {auth.create_token(dataset_session.get(AppUser, account), demo_role='sales')}"}
        for account in ("S2", "S3")
    }
    uploaded = api.post("/api/reports/datasets", headers=headers["S2"], files={"file": ("個人.csv", CSV.encode())})
    assert uploaded.status_code == 201
    dataset_id = uploaded.json()["id"]
    assert [item["id"] for item in api.get("/api/reports/datasets", headers=headers["S2"]).json()["datasets"]] == [dataset_id]
    for other in (headers["S3"], client(dataset_session, "S1")[1]):
        assert api.get("/api/reports/datasets", headers=other).json()["datasets"] == []
        assert api.get(f"/api/reports/datasets/{dataset_id}", headers=other).status_code == 404


def test_csv_infers_types_preserves_identifiers_and_complete_statistics():
    source, origin = service.parse_file(CSV.encode("utf-8-sig"), "銷售.csv")
    assert source["columns"][0]["kind"] == "text"
    assert source["columns"][2]["kind"] == "date"
    assert source["columns"][3]["kind"] == "number"
    columns, rows, changes = service.transform(source, service.default_recipe(source))
    assert rows[0] == ["001", "甲店", "2031-04-01", 1200, "正常"]
    assert service.profile(columns, rows)[3]["sum"] == 1400
    assert service.profile(columns, rows)[3]["empty"] == 1
    assert service.profile(columns, rows)[0]["min_length"] == service.profile(columns, rows)[0]["max_length"] == 3
    assert changes["can_save"] and origin["encoding"] == "utf-8-sig"


@pytest.mark.parametrize("filename,content,expected", [
    ("台灣.csv", "客戶,金額\n台北店,100\n".encode("cp950"), "cp950"),
    ("test.tsv", "客戶\t金額\n甲店\t100\n".encode(), "utf-8-sig"),
    ("test.csv", "客戶;金額\n甲店;100\n".encode("utf-16"), "utf-16"),
])
def test_encodings_and_delimiters(filename, content, expected):
    source, origin = service.parse_file(content, filename)
    assert len(source["columns"]) == 2 and len(source["rows"]) == 1
    assert origin["encoding"] == expected


def workbook():
    book = Workbook()
    book.active.title = "說明"
    book.active.append(["檔案說明"])
    book.active.append(["銷售資料在第二張表"])
    sheet = book.create_sheet("明細")
    sheet.append(["內部報表"])
    sheet.append(["編號", "金額", "公式"])
    sheet.append(["0001", 10, "=2+2"])
    raw = io.BytesIO(); book.save(raw)
    return raw.getvalue()


def test_excel_sheet_header_and_formula_cached_values():
    source, origin = service.parse_file(workbook(), "test.xlsx", "明細", 2)
    assert origin["sheets"] == ["說明", "明細"]
    assert source["rows"] == [["0001", 10, None]]
    assert origin["warnings"]
    with pytest.raises(service.DatasetError, match="工作表"):
        service.parse_file(workbook(), "test.xlsx", "不存在")


def test_excel_xml_entities_are_rejected():
    raw = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(workbook())) as original, zipfile.ZipFile(raw, "w") as changed:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                content = b'<!DOCTYPE workbook [<!ENTITY injected "blocked">]>' + content
            changed.writestr(entry.filename, content)
    with pytest.raises(service.DatasetError): service.parse_file(raw.getvalue(), "x.xlsx")


@pytest.mark.parametrize("content,name", [(b"", "x.csv"), (b"bad", "x.xlsx"), (b"PK\0\0", "x.csv"), (b"x,y\n", "x.csv"), (b"hello", "x.exe"), (b"a,b\n\"x", "x.csv")])
def test_bad_files_are_rejected(content, name):
    with pytest.raises(service.DatasetError): service.parse_file(content, name)


def test_limits_are_checked_without_silent_truncation(monkeypatch):
    monkeypatch.setattr(service, "MAX_ROWS", 2)
    with pytest.raises(service.DatasetError, match="最多"):
        service.parse_file(CSV.encode(), "x.csv")
    monkeypatch.setattr(service, "MAX_BYTES", 3)
    with pytest.raises(service.DatasetError): service.parse_file(b"abcd", "x.csv")
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"0" * (51 * 1024 * 1024))
    monkeypatch.setattr(service, "MAX_BYTES", 10 * 1024 * 1024)
    with pytest.raises(service.DatasetError, match="解壓"):
        service.parse_file(raw.getvalue(), "bomb.xlsx")


def test_cleaning_preview_is_explicit_and_original_rows_can_be_restored():
    source, _ = service.parse_file(b"id,value\n001,1\n001,1\n002,bad\n,\n", "x.csv")
    recipe = service.default_recipe(source)
    recipe.columns[1].kind = "number"
    recipe.deduplicate = True
    columns, rows, changes = service.transform(source, recipe)
    assert changes["invalid_cells"] == 1 and not changes["can_save"]
    assert changes["removed_blank"] == changes["removed_duplicates"] == 1
    assert rows == [["001", 1], ["002", None]]
    assert changes["issues"][0]["value"] == "bad"
    recipe.invalid = "null"
    assert service.transform(source, recipe)[2]["can_save"]
    assert source["rows"][2][1] == "bad"
    with pytest.raises(service.DatasetError): service.transform(source, service.Recipe(columns=[service.ColumnRule(id="c1", name="id")]))


def test_upload_roundtrip_query_full_totals_filters_and_export(dataset_session):
    api, headers, source = upload(dataset_session)
    id = source["id"]
    assert source["row_count"] == 4 and source["preview"]["profiles"][3]["sum"] == 1400
    assert api.get("/api/reports/datasets", headers=headers).json()["datasets"][0]["id"] == id
    result = api.post(f"/api/reports/datasets/{id}/query", headers=headers, json=chart()).json()
    assert result["total"] == 1400 and result["row_count"] == 4
    assert result["points"][0]["label"] == "甲店" and result["points"][1]["value"] == 200
    filtered = api.post(f"/api/reports/datasets/{id}/query", headers=headers, json=chart(aggregation="avg", filters=[{"column": "c2", "operator": "eq", "value": "乙店"}])).json()
    assert filtered["total"] == 100 and filtered["row_count"] == 2
    assert api.post(f"/api/reports/datasets/{id}/query", headers=headers, json=chart(value="c2")).status_code == 422
    assert api.post(f"/api/reports/datasets/{id}/query", headers=headers, json=chart(filters=[{"column": "c2", "operator": "gte", "value": "2"}])).status_code == 422
    assert api.get(f"/api/reports/datasets/{id}/export", headers=headers).text.startswith("\ufeff編號,客戶")
    assert api.get(f"/api/reports/datasets/{id}?page=2&page_size=2", headers=headers).json()["preview"]["rows"][0][0] == "003"
    assert dataset_session.scalar(select(func.count()).select_from(SalesTransaction)) == 6


def test_private_ownership_including_managers_and_all_mutations(dataset_session):
    api, _, source = upload(dataset_session)
    id = source["id"]
    for user in ("S2", "M1"):
        _, other = client(dataset_session, user)
        assert api.get("/api/reports/datasets", headers=other).json()["datasets"] == []
        for method, suffix, body in [("get", "", None), ("get", "/export", None), ("post", "/preview", source["recipe"]), ("post", "/query", chart()), ("post", "/analyze", {}), ("put", "", {"revision": 1, "name": "x", "recipe": source["recipe"], "charts": []})]:
            response = getattr(api, method)(f"/api/reports/datasets/{id}{suffix}", headers=other, **({"json": body} if body is not None else {}))
            assert response.status_code == 404, response.text
        assert api.delete(f"/api/reports/datasets/{id}?revision=1", headers=other).status_code == 404
    assert api.get("/api/reports/datasets").status_code == 401


def test_save_conflict_invalid_conversion_and_immutable_source(dataset_session):
    api, headers, source = upload(dataset_session)
    id = source["id"]
    recipe = source["recipe"]; recipe["columns"][1]["kind"] = "number"
    body = {"revision": 1, "name": "整理後", "recipe": recipe, "charts": []}
    assert api.post(f"/api/reports/datasets/{id}/preview", headers=headers, json=recipe).json()["changes"]["invalid_cells"] == 4
    assert api.put(f"/api/reports/datasets/{id}", headers=headers, json=body).status_code == 422
    recipe["invalid"] = "null"
    saved = api.put(f"/api/reports/datasets/{id}", headers=headers, json=body)
    assert saved.status_code == 200 and saved.json()["revision"] == 2
    assert api.put(f"/api/reports/datasets/{id}", headers=headers, json=body).status_code == 409
    restore = saved.json()["recipe"]; restore["columns"][1]["kind"] = "text"
    restored = api.put(f"/api/reports/datasets/{id}", headers=headers, json={**body, "revision": 2, "recipe": restore}).json()
    assert restored["preview"]["rows"][0][1] == "甲店"


def test_save_chart_persists_and_reparse_resets_it(dataset_session):
    api, headers, source = upload(dataset_session, workbook(), "test.xlsx", sheet="明細", header_row=2)
    id = source["id"]
    saved = api.put(f"/api/reports/datasets/{id}", headers=headers, json={"revision": 1, "name": source["name"], "recipe": source["recipe"], "charts": [chart(group_by="c1", value="c2")]}).json()
    assert len(saved["charts"]) == 1
    response = api.post(f"/api/reports/datasets/{id}/reparse", headers=headers, json={"revision": 2, "sheet": "說明", "header_row": 1})
    assert response.status_code == 200 and response.json()["charts"] == []
    assert response.json()["preview"]["rows"][0] == ["銷售資料在第二張表"]


def test_append_and_left_join_preserve_sources_and_reject_duplicate_keys(dataset_session):
    api, headers, left = upload(dataset_session, "id,value\n001,10\n002,20\n")
    _, _, right = upload(dataset_session, "id,label\n001,甲\n003,丙\n")
    body = {"left_id": left["id"], "right_id": right["id"], "left_revision": 1, "right_revision": 1, "name": "合併", "mode": "append"}
    appended = api.post("/api/reports/datasets/merge", headers=headers, json=body)
    assert appended.status_code == 201, appended.text
    assert appended.json()["row_count"] == 4 and len(appended.json()["columns"]) == 3
    joined = api.post("/api/reports/datasets/merge", headers=headers, json={**body, "mode": "left_join", "left_key": "c1", "right_key": "c1"})
    assert joined.status_code == 201, joined.text
    assert joined.json()["preview"]["rows"] == [["001", 10, "001", "甲"], ["002", 20, None, None]]
    assert api.get(f"/api/reports/datasets/{left['id']}", headers=headers).json()["revision"] == 1
    assert api.post("/api/reports/datasets/merge", headers=headers, json={**body, "left_revision": 99}).status_code == 409
    _, _, duplicates = upload(dataset_session, "id,label\n001,甲\n001,乙\n")
    assert api.post("/api/reports/datasets/merge", headers=headers, json={**body, "right_id": duplicates["id"], "mode": "left_join", "left_key": "c1", "right_key": "c1"}).status_code == 422
    _, other = client(dataset_session, "S2")
    assert api.post("/api/reports/datasets/merge", headers=other, json=body).status_code == 404


def test_quota_delete_revision_and_csv_formula_safety(dataset_session, monkeypatch):
    api, headers, source = upload(dataset_session, "name,value\n=HYPERLINK(1),-20\n")
    assert "'=HYPERLINK(1),-20" in api.get(f"/api/reports/datasets/{source['id']}/export", headers=headers).text
    monkeypatch.setattr(service, "MAX_DATASETS", 1)
    assert api.post("/api/reports/datasets", headers=headers, files={"file": ("x.csv", CSV)}).status_code == 409
    assert api.delete(f"/api/reports/datasets/{source['id']}?revision=2", headers=headers).status_code == 409
    assert api.delete(f"/api/reports/datasets/{source['id']}?revision=1", headers=headers).status_code == 204
    assert api.get(f"/api/reports/datasets/{source['id']}", headers=headers).status_code == 404


def test_ai_receives_real_rows_full_statistics_and_validates_unknown_columns():
    raw, _ = service.parse_file(CSV.encode(), "x.csv")
    columns, rows, _ = service.transform(raw, service.default_recipe(raw))
    dataset = SimpleNamespace(columns=columns, rows=rows, name="銷售", origin={"filename": "x.csv"}, revision=3)
    class Model:
        result = {"summary": "這份檔案是銷售資料，金額合計 1400。", "findings": ["金額有一格空白"], "columns": [], "charts": [chart(id="model_chart_1")]}
        def json(self, **kwargs):
            prompt = json.loads(kwargs["prompt"])
            assert prompt["statistics"][3]["sum"] == 1400
            assert prompt["sample_rows"][0]["values"][0] == "001"
            assert prompt["row_count"] == 4 and "不可信資料" in kwargs["system"]
            return self.result
    model = Model()
    analysis = service.analyze(dataset, "有什麼異常？", model)
    assert analysis["sample_rows"] == analysis["total_rows"] == 4 and analysis["revision"] == 3
    assert analysis["charts"][0]["id"].startswith("chart-")
    model.result = {"summary": "測試", "columns": [{"id": "unknown", "name": "未知", "kind": "text", "reason": "測試"}]}
    with pytest.raises(LLMOutputError): service.analyze(dataset, "看檔案", model)


def test_ai_unavailable_is_clear_and_analyze_has_no_usage_limit(dataset_session):
    api, headers, source = upload(dataset_session)
    response = api.post(f"/api/reports/datasets/{source['id']}/analyze", headers=headers, json={})
    assert response.status_code == 503 and "AI 尚未設定" in response.json()["detail"]
    assert usage.buckets_for("POST", f"/api/reports/datasets/{source['id']}/analyze") == []


def test_inspect_does_not_persist_and_empty_excel_cover_can_select_another_sheet(dataset_session):
    api, headers = client(dataset_session)
    book = Workbook()
    sheet = book.create_sheet("資料")
    sheet.append(["id", "金額"]); sheet.append(["001", 100])
    raw = io.BytesIO(); book.save(raw)
    initial = api.post("/api/reports/datasets/inspect", headers=headers, files={"file": ("x.xlsx", raw.getvalue())})
    assert initial.status_code == 200
    assert initial.json()["origin"]["sheets"] == ["Sheet", "資料"]
    assert initial.json()["preview"] is None
    selected = api.post("/api/reports/datasets/inspect", headers=headers, files={"file": ("x.xlsx", raw.getvalue())}, data={"sheet": "資料"})
    assert selected.status_code == 200 and selected.json()["preview"]["rows"] == [["001", 100]]
    assert dataset_session.scalar(select(func.count()).select_from(ReportDataset)) == 0


def test_aggregation_uses_all_rows_and_pie_has_other_and_rejects_non_additive_values():
    source, _ = service.parse_file("名稱,金額\nA,10\nB,20\nC,30\nD,40\nE,50\n".encode(), "x.csv")
    columns, rows, _ = service.transform(source, service.default_recipe(source))
    config = service.DatasetChart(**chart(group_by="c1", value="c2", limit=3))
    result = service.aggregate(columns, rows, config)
    assert result["total"] == 150 and len(result["points"]) == 3
    pie = service.aggregate(columns, rows, config.model_copy(update={"visual": "pie"}))
    assert len(pie["points"]) == 4 and sum(point["value"] for point in pie["points"]) == 150
    with pytest.raises(service.DatasetError): service.aggregate(columns, rows, config.model_copy(update={"visual": "pie", "aggregation": "avg"}))


def test_ai_api_uses_saved_owner_data_and_returns_structured_validated_proposals(dataset_session, monkeypatch):
    from app.api import report_datasets as api_module
    api, headers, source = upload(dataset_session)
    monkeypatch.setattr(api_module, "ai_available", lambda: True)
    class Model:
        def json(self, **kwargs):
            content = json.loads(kwargs["prompt"])
            assert content["statistics"][3]["sum"] == 1400
            assert content["sample_rows"][0]["values"][0] == "001"
            return {"summary": "讀到銷售資料，共四列。", "findings": ["金額合計為 1400。"], "cautions": [], "columns": [], "charts": [chart()]}
    monkeypatch.setattr(service, "get_llm", lambda: Model())
    result = api.post(f"/api/reports/datasets/{source['id']}/analyze", headers=headers, json={"question": "請看實際資料"})
    assert result.status_code == 200, result.text
    assert result.json()["revision"] == 1 and result.json()["charts"][0]["value"] == "c4"
    assert api.get(f"/api/reports/datasets/{source['id']}", headers=headers).json()["charts"] == []


def test_imported_source_and_original_file_restore_with_account_settings(dataset_session):
    import seed
    _, _, source = upload(dataset_session)
    dataset = dataset_session.get(ReportDataset, source["id"])
    snapshot = {column.key: getattr(dataset, column.key) for column in ReportDataset.__table__.columns}
    dataset_session.delete(dataset); dataset_session.flush()
    kept = seed.KeptAccounts(users=[], identities=[], settings={"report_dataset": [snapshot]})
    assert seed.restore_settings(dataset_session, kept)["report_dataset"] == 1
    restored = dataset_session.get(ReportDataset, source["id"])
    assert restored.content == CSV.encode()
    assert restored.rows[0][0] == "001" and restored.row_count == 4
