import csv
import datetime as dt
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from app.main import app
from app.models import AppUser, Customer, Product, Receivable, SalesTransaction, Visit
from app.services import auth, report_exports as exports, reports


def config(card=None, **patch):
    return reports.ReportQuery.model_validate({"name": "完整匯出驗收", "template": "manager", "period": "this_month", "cards": [{"id": "sales", "title": "銷售完整分組", "metric": "sales", "dimension": "customer", "visual": "bar", "limit": 3, **(card or {})}], **patch})


def body(card=None, *, format="csv", data="summary", query=None, **patch):
    return exports.ExportInput(config=query or config(card), format=format, data=data, **patch)


def table(session, request, user="M1"):
    return exports.tables(session, session.get(AppUser, user), request)[0]


def root(records):
    return next(item for item in records if item["_record"] == "總計" and item["row_label"] == item["column_label"] == "總計")


def transaction(session, order, customer="C1", sku="P1", amount=1, cost=0, date="2031-04-01", fee=0, reward=0):
    session.add(SalesTransaction(order_no=order, customer_id=customer, sku=sku, date=dt.date.fromisoformat(date), qty=1, amount=amount, cost=cost, listing_fee=fee, channel_reward=reward))


def analytical_fixture(session):
    session.add(Product(sku="P0", name="零金額品項", category="測試", spec="盒", unit="盒", unit_price=0, unit_cost=10))
    session.flush()
    transaction(session, "O1", sku="P2", amount=100, fee=20, reward=10)
    transaction(session, "O1", sku="P0", amount=0, cost=10)
    session.flush()


def matrix():
    return {"visual": "matrix", "dimension": "owner", "legend": "year", "matrix": {"rows": ["owner", "customer_type", "customer"], "columns": ["year", "month", "product"], "metrics": ["sales", "net_margin_rate", "order_count", "avg_order_value", "net_margin"]}}


def test_complete_groups_ignore_top_n_and_csv_carries_metadata_and_total(interaction_session):
    for index in range(26):
        key = f"E{index}"
        interaction_session.add(Customer(id=key, name=key, type="clinic", region="北區", city="臺北", place_id="TPE", grade="A", owner_user_id="S1", area="中正", lat=25, lng=121, payment_term="RM06"))
        transaction(interaction_session, key, customer=key)
    interaction_session.flush()
    request = body()
    assert len(reports.data(interaction_session, interaction_session.get(AppUser, "M1"), request.config).cards[0].points) == 3
    complete = table(interaction_session, request)
    records = list(complete.records())
    assert len([row for row in records if row["_record"] == "最低層分組"]) == 29
    assert root(records)["sales"] == 1376
    assert all("其他" not in row["row_label"] for row in records)
    stream = exports.csv_file(request, [complete])
    content = stream.read(); stream.close()
    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    assert len(rows) == 31 and rows[0]["資料類型"] == "匯出條件"
    metadata = json.loads(rows[0]["範圍與條件"])
    assert metadata["complete"] is True and metadata["total"] == 1376
    assert rows[-1]["資料類型"] == "總計"


def test_matrix_complete_all_levels_is_independent_of_views_and_recalculates_distinct_and_rates(interaction_session):
    analytical_fixture(interaction_session)
    query = config(matrix())
    first = list(table(interaction_session, body(query=query)).records())
    query.matrix_views = {"sales": reports.MatrixView(row_page=2, column_page=2, row_page_size=25, column_page_size=5, row_expand_depth=1, column_expand_depth=1)}
    assert list(table(interaction_session, body(query=query)).records()) == first
    total = root(first)
    assert total["sales"] == 1450 and total["net_margin"] == 735
    assert total["net_margin_rate"] == pytest.approx(735 / 1450)
    assert total["order_count"] == 3 and total["avg_order_value"] == pytest.approx(1450 / 3)
    zero = next(row for row in first if row["_record"] == "最低層分組" and "C1" in row["row_label"] and "零金額品項" in row["column_label"])
    assert zero["net_margin"] == -10 and zero["net_margin_rate"] is None and zero["avg_order_value"] == 0 and zero["order_count"] == 1
    assert not any("C4" in row["row_label"] and row["column_label"].endswith("P1") and row["_record"] == "最低層分組" for row in first)
    assert any(row["_record"] == "小計" for row in first)


def test_full_comparisons_calendar_gaps_and_nulls(interaction_session):
    transaction(interaction_session, "PREVIOUS", amount=200, cost=100, date="2031-03-17")
    interaction_session.flush()
    request = body({"dimension": "month", "visual": "line", "comparison": "previous_month"})
    records = list(table(interaction_session, request).records())
    total = root(records)
    assert total["sales_previous"] == 200 and total["sales_delta"] == 1150
    assert total["sales_growth"] == pytest.approx(5.75)
    assert next(row for row in records if row["row_label"] == "2031/04")["sales_previous"] == 200
    ratio = config({"dimension": "month", "visual": "line", "metric": "net_margin_rate", "comparison": "previous_month"}, period="six_months")
    records = list(table(interaction_session, body(query=ratio)).records())
    february = next(row for row in records if row["row_label"] == "2031/02")
    assert february["net_margin_rate"] is None and february["row_count"] == 0
    january = next(row for row in records if row["row_label"] == "2031/01")
    assert january["net_margin_rate_previous"] is None
    assert january["net_margin_rate_status"] == "上期無資料"


def test_export_interaction_modes_respect_fixed_filters_and_scope(interaction_session):
    query = config({"dimension": "owner"})
    query.cards += [reports.ReportCard(id="target", title="接受圖", metric="sales", dimension="customer", visual="bar", interaction="highlight")]
    query.selection_groups = [reports.ReportSelectionGroup(source_id="sales", points=[reports.ReportSelectionPoint(filters=[reports.ReportFilter(dimension="owner", key="S1")], label="S1")])]
    full = exports.tables(interaction_session, interaction_session.get(AppUser, "M1"), body(query=query))
    subset = exports.tables(interaction_session, interaction_session.get(AppUser, "M1"), body(query=query, selection_scope="selected"))
    assert full[1].total == 1350 and subset[1].total == 650 and subset[0].total == 1350
    query.cards[1].interaction = "none"
    query.slicers = [reports.ReportSlicer(dimension="product", keys=["P1"])]
    assert exports.tables(interaction_session, interaction_session.get(AppUser, "M1"), body(query=query, selection_scope="selected"))[1].total == 1100
    assert table(interaction_session, body(), "S1").total == 650
    private = config(slicers=[{"dimension": "customer", "keys": ["C3"]}])
    assert table(interaction_session, body(query=private)).total == 0


def test_full_raw_csv_and_xlsx_keep_all_rows_types_dates_and_formula_protection(interaction_session):
    for index in range(124):
        transaction(interaction_session, f"RAW-{index:03}")
    interaction_session.get(Customer, "C1").name = '=HYPERLINK("https://example.invalid","客戶")\n第二行'
    interaction_session.flush()
    request = body(data="details", format="xlsx")
    exported = table(interaction_session, request)
    assert exported.count == 127
    stream, filename = exports.export_file(interaction_session, interaction_session.get(AppUser, "M1"), request)
    book = load_workbook(stream, data_only=False)
    sheet = book.worksheets[1]
    rows = list(sheet.iter_rows())
    headers = {cell.value: index for index, cell in enumerate(rows[4])}
    assert len(rows[5:]) == 127 and filename.endswith(".xlsx")
    customer_cells = [row[headers["客戶"]] for row in rows[5:] if row[headers["客戶編號"]].value == "C1"]
    assert customer_cells and all(cell.data_type == "s" and cell.value.startswith("'=") for cell in customer_cells)
    assert all(isinstance(row[headers["交易日期"]].value, dt.datetime) for row in rows[5:])
    assert all(row[headers["業績（元）"]].data_type == "n" for row in rows[5:])
    assert sheet.freeze_panes == "A6" and sheet.auto_filter.ref.endswith("132")
    stream.close()
    csv_stream = exports.csv_file(request, [exported])
    csv_rows = list(csv.DictReader(io.StringIO(csv_stream.read().decode("utf-8-sig")))); csv_stream.close()
    assert len([row for row in csv_rows if row["資料類型"] == "實際明細"]) == 127
    assert any(row["客戶"].startswith("'=") and "\n第二行" in row["客戶"] for row in csv_rows)


def test_matrix_six_condition_point_xlsx_keeps_only_actual_line_and_selected_metric(interaction_session):
    analytical_fixture(interaction_session)
    query = config(matrix())
    point = reports.ReportSelectionPoint(filters=[reports.ReportFilter(dimension=dimension, key=key) for dimension, key in [("owner", "S1"), ("customer_type", "clinic"), ("customer", "C1"), ("year", "2031"), ("month", "2031-04"), ("product", "P2")]], label="C1 / P2", metric="net_margin_rate")
    exported = table(interaction_session, body(query=query, data="details", format="xlsx", card_id="sales", point=point))
    assert exported.count == 1 and exported.total == pytest.approx(.7)
    records = list(exported.records())
    assert records[0]["order"] == "O1" and records[0]["amount"] == 100 and records[0]["net_margin"] == 70
    assert exported.metrics == ["net_margin_rate"]


def test_xlsx_rate_style_and_duplicate_long_sheet_names(interaction_session):
    query = config({"metric": "net_margin_rate", "format": {"numbers": {"net_margin": {"decimals": 6}, "net_margin_rate": {"decimals": 1}}}})
    query.cards[0].title = "中文名稱/" * 7
    query.cards.append(query.cards[0].model_copy(update={"id": "second"}))
    # title 限制在 40 字；Excel 工作表需要另外截至 31 並去除保留字元。
    query.cards[0].title = query.cards[1].title = "中文名稱/" * 6
    stream, _ = exports.export_file(interaction_session, interaction_session.get(AppUser, "M1"), body(query=query, format="xlsx"))
    book = load_workbook(stream)
    assert len({name.casefold() for name in book.sheetnames}) == 3 and all(len(name) <= 31 and "/" not in name for name in book.sheetnames)
    sheet = book.worksheets[1]
    header = [cell.value for cell in sheet[5]]
    column = header.index("淨毛利率（比例）") + 1
    assert sheet.cell(6, column).value == .5 and sheet.cell(6, column).number_format == "#,##0.0%"
    stream.close()


def test_current_snapshots_and_taipei_visits_keep_their_semantics(interaction_session):
    interaction_session.add(Visit(id="EXPORT-VISIT", customer_id="C1", user_id="S2", visited_at=dt.datetime.fromisoformat("2031-04-09T16:00:00+00:00"), status="confirmed"))
    interaction_session.add(Receivable(invoice_no="EXPORT-OLD", customer_id="C1", invoice_date=dt.date(2029, 1, 1), due_date=dt.date(2029, 2, 1), paid_date=None, amount=100))
    interaction_session.flush()
    visits = body({"metric": "visits", "dimension": "owner"}, data="details", format="xlsx")
    stream, _ = exports.export_file(interaction_session, interaction_session.get(AppUser, "M1"), visits)
    book = load_workbook(stream)
    rows = list(book.worksheets[1].values)
    headers = list(rows[4])
    assert rows[5][headers.index("實際拜訪時間（台北）")] == dt.datetime(2031, 4, 10, 0, 0)
    assert rows[5][headers.index("實際拜訪人員")] == "S2"
    stream.close()
    overdue = table(interaction_session, body({"metric": "overdue"}))
    assert overdue.total == 100 and "快照" in overdue.note


def test_empty_date_intersection_preserves_null_in_complete_export(interaction_session):
    query = config({"metric": "net_margin_rate"}, pages=[{"id": "overview", "name": "無日期交集", "date_range": {"start": "2031-03-01", "end": "2031-03-31"}}])
    request = body(query=query)
    exported = table(interaction_session, request)
    records = list(exported.records())
    assert len(records) == 1 and root(records)["net_margin_rate"] is None and exported.count == 0
    assert "日期沒有交集" in exported.note
    stream = exports.csv_file(request, [exported])
    rows = list(csv.DictReader(io.StringIO(stream.read().decode("utf-8-sig")))); stream.close()
    assert rows[-1]["淨毛利率（比例）"] == ""


def test_pdf_contains_all_pages_hidden_mobile_cards_and_embedded_chinese_font(interaction_session):
    query = config()
    query.pages.append(reports.ReportPage(id="analysis", name="完整分析頁"))
    query.cards.append(reports.ReportCard(id="customers", title="手機隱藏的客戶數", metric="customers", visual="kpi", dimension="none", page_id="analysis", mobile_layout=reports.MobileLayout(order=0, width="full", height=8, hidden=True)))
    stream, filename = exports.export_file(interaction_session, interaction_session.get(AppUser, "M1"), body(query=query, format="pdf"))
    reader = PdfReader(stream)
    text = "".join(page.extract_text() for page in reader.pages)
    assert len(reader.pages) >= 3 and "完整分析頁" in text and "手機隱藏的客戶數" in text
    assert "1,350.00 元" in text and "完整彙總" in text and "第 1 頁" in text
    fonts = [font.get_object() for page in reader.pages for font in page["/Resources"].get("/Font", {}).get_object().values()]
    assert any("/FontFile2" in font.get("/FontDescriptor", {}).get_object() for font in fonts if "/FontDescriptor" in font)
    assert filename.endswith(".pdf")
    stream.close()


@pytest.mark.parametrize("patch", [{"format": "zip"}, {"format": "pdf", "data": "details"}, {"format": "pdf", "page_id": "overview"}, {"card_id": "missing"}, {"page_id": "missing"}, {"page_id": "overview", "card_id": "sales"}])
def test_invalid_export_requests_are_rejected(patch):
    with pytest.raises(ValidationError):
        exports.ExportInput.model_validate({"config": config().model_dump(mode="json"), "format": "csv", **patch})


def test_export_api_enforces_ownership_and_download_headers(interaction_session):
    client = TestClient(app)
    headers = {"Authorization": f"Bearer {auth.create_token(interaction_session.get(AppUser, 'S1'))}"}
    response = client.post("/api/reports/export", headers=headers, json=body().model_dump(mode="json"))
    assert response.status_code == 200 and response.content.startswith(b"\xef\xbb\xbf")
    assert "filename*=UTF-8''" in response.headers["content-disposition"] and response.headers["cache-control"] == "no-store"
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert not any("C2" in row.get("列階層", "") or "C3" in row.get("列階層", "") for row in rows)
    private = body(query=config(owner_id="S3"))
    assert client.post("/api/reports/export", headers=headers, json=private.model_dump(mode="json")).status_code == 403
    invalid = body().model_dump(mode="json"); invalid["page"] = 2
    assert client.post("/api/reports/export", headers=headers, json=invalid).status_code == 422
