"""可攜式 Excel 與含嵌入中文字型的完整報表 PDF。"""

from __future__ import annotations

import datetime as dt
import math
import re
import threading
from decimal import Decimal
from pathlib import Path
from tempfile import SpooledTemporaryFile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Flowable, KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app.services import reports as r
from app.services.report_exports import DIMENSIONS, METRICS, ExportInput, ExportTable, json_text, safe_cell

FONT = "ReportHuninn"
FONT_LOCK = threading.Lock()
INK = colors.HexColor("#273044")
PURPLE = colors.HexColor("#6455c7")
PALE = colors.HexColor("#f3f1fb")


def pdf_text(value):
    mapping = pdfmetrics.getFont(FONT).face.charToGlyph
    return "".join(character if character in "\n\t" or ord(character) in mapping else f"\\u{ord(character):04x}" for character in str(value))


def _number_format(card, key, kind):
    metric = key if key in METRICS else next((metric for metric in sorted(METRICS, key=len, reverse=True) if key.startswith(metric + "_")), card.metric)
    decimals = card.format.numbers.get(metric)
    places = decimals.decimals if decimals and decimals.decimals is not None else 0 if kind == "number" and not key.endswith("_delta") else 2
    if key.endswith("_growth"): places = 2
    return ("#,##0" + ("." + "0" * places if places else "") + ("%" if kind == "percent" else ""))


def excel_cell(value):
    if isinstance(value, dt.datetime): return value.astimezone(r.TAIPEI).replace(tzinfo=None)
    if isinstance(value, dt.date): return value
    value = safe_cell(value)
    return ILLEGAL_CHARACTERS_RE.sub(lambda match: f"\\u{ord(match[0]):04x}", value) if isinstance(value, str) else value


def xlsx_file(body: ExportInput, tables: list[ExportTable]):
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    book = Workbook(write_only=True)
    book.properties.title = excel_cell(body.config.name)
    book.properties.creator = "報表工作區"
    book.properties.subject = "權限內完整報表資料；各圖表依自己的有效條件計算"
    used = set()
    def sheet_name(title):
        title = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "_", title).strip("'") or "報表"
        result, counter = title[:31], 1
        while result.casefold() in used:
            counter += 1; suffix = f"_{counter}"; result = title[:31-len(suffix)] + suffix
        used.add(result.casefold()); return result
    index = book.create_sheet(sheet_name("匯出說明"))
    index.column_dimensions["A"].width = 26
    index.column_dimensions["B"].width = 110
    index_row = 0
    def info(label, value):
        nonlocal index_row
        text = str(value)
        pieces = [text[start:start+800] for start in range(0, len(text), 800)] or [""]
        for part, piece in enumerate(pieces):
            index_row += 1
            left = WriteOnlyCell(index, excel_cell(label if part == 0 else label + "（續）")); left.font = Font(bold=True, color="6455C7"); left.alignment = Alignment(vertical="top")
            right = WriteOnlyCell(index, excel_cell(piece)); right.alignment = Alignment(wrap_text=True, vertical="top")
            index.row_dimensions[index_row].height = min(310, max(24, math.ceil(len(piece) / 50) * 15))
            index.append([left, right])
    info("報表", body.config.name)
    info("匯出時間（台北）", dt.datetime.now(r.TAIPEI).isoformat(timespec="seconds"))
    info("資料內容", "完整彙總" if body.data == "summary" else "完整實際明細")
    info("讀取方式", "每個圖表獨立套用報表、分頁、單圖及探索條件；不同圖表可能包含相同紀錄，不能跨工作表直接相加。數值保留原始單位，百分比儲存 0-1 並以百分比顯示，百分點差額使用百分點。")
    info("資料範圍", "全部符合目前資料權限的紀錄，不受畫面前 N 筆與分頁限制。彙總含完整實際分組與各層小計／總計；矩陣不補沒有交易的客戶／商品配對。明細只含本期實際紀錄。")
    for table in tables:
        info("圖表", f"{table.page.name} / {table.card.title} ({table.card.id})")
        info("權限與日期", f"{table.scope}；資料基準日 {table.as_of}；本期 {table.start} 至 {table.end}")
        info("定義", table.note)
        info("完整條件 JSON", json_text(table.conditions))
        part = 0
        count = 0
        sheet = None
        def new_sheet():
            nonlocal part, count
            part += 1; count = 5
            title = sheet_name(f"{table.page.name}-{table.card.title}" + (f"_{part}" if part > 1 else ""))
            sheet = book.create_sheet(title)
            sheet.freeze_panes = "A6"
            sheet.sheet_view.showGridLines = False
            sheet.column_dimensions["A"].width = 18
            sheet.row_dimensions[1].height = 42
            sheet.row_dimensions[2].height = 48
            sheet.row_dimensions[4].height = 36
            for number, (key, _, kind) in enumerate(table.columns, start=2):
                dim = sheet.column_dimensions[get_column_letter(number)]
                dim.width = 42 if key in ("row_label", "column_label", "product") else 30 if kind == "text" else 22
                if key in ("row_filters", "column_filters"):
                    dim.hidden = True
            title_cell = WriteOnlyCell(sheet, excel_cell(table.card.title)); title_cell.alignment = Alignment(wrap_text=True, vertical="top"); title_cell.font = Font(bold=True, size=14, color="6455C7")
            page_cell = WriteOnlyCell(sheet, excel_cell(table.page.name)); page_cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.append(["圖表", title_cell])
            sheet.append(["分頁", page_cell, "資料範圍", excel_cell(table.scope)])
            total_cell = WriteOnlyCell(sheet, table.total); total_cell.number_format = _number_format(table.card, table.card.metric, "percent" if table.card.metric == "net_margin_rate" else "currency" if table.card.metric in ("sales", "net_margin", "avg_order_value", "overdue") else "number")
            sheet.append(["完整範圍合計", total_cell, "符合紀錄數", table.count])
            sheet.append(["有效條件及定義", "參見「匯出說明」；編號與條件欄預設隱藏，可在 Excel 展開。"])
            headers = []
            for label in ["資料類型", *(label for _, label, _ in table.columns)]:
                cell = WriteOnlyCell(sheet, label); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="6455C7"); cell.alignment = Alignment(wrap_text=True, vertical="center"); headers.append(cell)
            sheet.row_dimensions[5].height = 34
            sheet.append(headers)
            info("資料工作表", title)
            return sheet
        for record in table.records():
            if sheet is None or count >= 1048576:
                if sheet is not None: sheet.auto_filter.ref = f"A5:{get_column_letter(len(table.columns)+1)}{count}"
                sheet = new_sheet()
            row = [record.get("_record", "實際明細")]
            for key, _, kind in table.columns:
                value = record.get(key)
                cell = WriteOnlyCell(sheet, excel_cell(value) if value is not None else None)
                if kind in ("number", "currency", "percent") and isinstance(value, (int, float, Decimal)):
                    cell.number_format = _number_format(table.card, key, kind)
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                elif kind in ("date", "datetime") and value is not None:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss" if kind == "datetime" else "yyyy-mm-dd"
                    cell.alignment = Alignment(vertical="top")
                else: cell.alignment = Alignment(vertical="top", wrap_text=True)
                if record.get("_record") in ("總計", "小計"):
                    cell.fill = PatternFill("solid", fgColor="F3F1FB"); cell.font = Font(bold=True, color="273044")
                row.append(cell)
            sheet.append(row); count += 1
        if sheet is None: sheet = new_sheet()
        sheet.auto_filter.ref = f"A5:{get_column_letter(len(table.columns)+1)}{count}"
    book.save(output); output.seek(0)
    return output


def exact(value, metric, card=None):
    if value is None: return "無法計算"
    percent = metric == "net_margin_rate"
    number = value * 100 if percent else value
    setting = card.format.numbers.get(metric) if card else None
    places = setting.decimals if setting and setting.decimals is not None else 0 if metric in ("order_count", "visits", "customers") else 2
    text = f"{number:,.{places}f}"
    return text + ("%" if percent else " 張" if metric == "order_count" else " 次" if metric == "visits" else " 家" if metric == "customers" else " 元")


class ValuePreview(Flowable):
    """列印用數值摘要，完整資料表緊接其後，不冒充畫布的像素截圖。"""
    def __init__(self, rows, metric, card, width):
        super().__init__(); self.rows = rows; self.metric = metric; self.card = card; self.width = width; self.height = 24 + len(rows) * 21

    def draw(self):
        canvas = self.canv
        values = [record[self.metric] for record in self.rows if record.get(self.metric) is not None]
        low, high = min([0, *values]), max([0, *values])
        span = high - low or 1
        left, width = 170, self.width - 310
        zero = left + (-low / span) * width
        canvas.setStrokeColor(colors.HexColor("#cbd1dd")); canvas.line(zero, 4, zero, self.height - 12)
        canvas.setFont(FONT, 8)
        for index, record in enumerate(self.rows):
            y = self.height - 25 - index * 21
            label = pdf_text(record["row_label"] + (" / " + record["column_label"] if record["column_label"] != "總計" else ""))
            while pdfmetrics.stringWidth(label, FONT, 8) > left - 12: label = label[:-2] + "…"
            canvas.setFillColor(INK); canvas.drawString(0, y + 3, label)
            value = record.get(self.metric)
            if value is not None:
                endpoint = left + (value - low) / span * width
                canvas.setFillColor(colors.HexColor("#d76857") if value < 0 else colors.HexColor(self.card.color))
                canvas.roundRect(min(zero, endpoint), y, max(abs(endpoint-zero), .6), 12, 2, stroke=0, fill=1)
            canvas.setFillColor(INK); canvas.drawRightString(self.width, y + 3, exact(value, self.metric, self.card))


def pdf_file(body: ExportInput, exported: list[ExportTable]):
    with FONT_LOCK:
        if FONT not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(FONT, str(Path(__file__).resolve().parents[1] / "assets/fonts/Huninn-Regular.ttf")))
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    page_size = landscape(A4)
    width = page_size[0] - 72
    styles = {
        "title": ParagraphStyle("ReportTitle", fontName=FONT, fontSize=24, leading=32, textColor=INK, spaceAfter=15),
        "page": ParagraphStyle("ReportPage", fontName=FONT, fontSize=18, leading=25, textColor=PURPLE, spaceAfter=12),
        "card": ParagraphStyle("ReportCard", fontName=FONT, fontSize=14, leading=20, textColor=INK, spaceAfter=7),
        "body": ParagraphStyle("ReportBody", fontName=FONT, fontSize=9, leading=14, textColor=INK, spaceAfter=5, wordWrap="CJK"),
        "small": ParagraphStyle("ReportSmall", fontName=FONT, fontSize=8, leading=11, textColor=INK, wordWrap="CJK"),
        "header": ParagraphStyle("ReportHead", fontName=FONT, fontSize=8, leading=11, textColor=colors.white, wordWrap="CJK"),
    }
    def paragraph(text, style="body"):
        return Paragraph(escape(pdf_text(text)).replace("\n", "<br/>"), styles[style])
    story = [paragraph(body.config.name, "title"), paragraph("完整報表 · 權限內彙總資料", "page"), paragraph(f"匯出時間（台北）：{dt.datetime.now(r.TAIPEI).isoformat(timespec='seconds')}")]
    if exported: story += [paragraph(f"資料基準日：{exported[0].as_of}；可見範圍：{exported[0].scope}")]
    story += [paragraph("此 PDF 重新排版所有分頁與圖表，包含桌機圖表及完整實際分組、小計與總計，不受畫面前 N 筆、矩陣分頁／展開或手機隱藏設定限制。數值摘要圖只預覽部分分組，後方表格保留全部資料。"),
        paragraph("各圖表依自己的有效條件及聯動模式計算，可能包含相同紀錄，不能跨圖表直接加總。去重訂單、毛利率與客單價直接重新聚合；沒有交易的客戶／商品配對不補虛構列，日曆空桶保留。"),
        paragraph("報表固定篩選：" + json_text({"period": body.config.period, "owner_id": body.config.owner_id, "customer_type": body.config.customer_type, "filters": body.config.filters, "slicers": body.config.slicers, "date_range": body.config.date_range}))]
    if body.config.selection_groups or body.config.drill_filters:
        story.append(paragraph("聯動與下鑽條件：" + json_text({"selection_scope": body.selection_scope, "selection_groups": body.config.selection_groups, "drill_filters": body.config.drill_filters})))
    story.append(PageBreak())
    for page_index, page in enumerate(body.config.pages):
        if page_index: story.append(PageBreak())
        story += [paragraph(page.name, "page"), paragraph("分頁篩選：" + json_text({"filters": page.filters, "slicers": page.slicers, "date_range": page.date_range}))]
        page_tables = [table for table in exported if table.page.id == page.id]
        if not page_tables: story.append(paragraph("本頁沒有圖表。"))
        for table in page_tables:
            data = list(table.records())
            heading = [paragraph(table.card.title, "card"), paragraph(f"資料基準日 {table.as_of}；" + ("目前快照" if table.card.metric in ("customers", "overdue") else f"本期 {table.start} 至 {table.end}") + f"；符合紀錄 {table.count:,} 筆")]
            total = next((record for record in data if record.get("_record") == "總計" and record.get("row_label") == record.get("column_label") == "總計"), {})
            heading.append(paragraph("完整範圍合計：" + "；".join(METRICS[metric].replace("（比例）", "（%）") + " " + exact(total.get(metric), metric, table.card) for metric in table.metrics)))
            story.append(KeepTogether(heading))
            story.append(paragraph(table.note))
            comparison_context = table.conditions.get("comparison", {})
            if comparison_context.get("mode") != "none" and comparison_context:
                dates = comparison_context.get("dates")
                story.append(paragraph(("前月同區間" if comparison_context["mode"] == "previous_month" else "去年同期") + (f"：{dates[0]} 至 {dates[1]}" if dates else "：比較日期超出可用範圍。")))
            story.append(paragraph("本圖條件：" + json_text({"filters": table.card.filters, "slicers": table.card.slicers, "date_range": table.card.date_range, "drill_filters": body.config.drill_filters.get(table.card.id, []), "interaction": table.card.interaction, "selection_scope": body.selection_scope})))
            leaf = [record for record in data if record.get("_record") == "最低層分組"]
            if table.card.visual not in ("matrix", "table", "kpi") and leaf:
                preview = leaf[:table.card.limit]
                story += [paragraph(f"數值摘要圖（前 {len(preview)} 個實際分組；完整資料如下）", "small"), ValuePreview(preview, table.card.metric, table.card, width), Spacer(1, 8)]
            for metric in table.metrics:
                comparison = any(key == metric + "_previous" for key, _, _ in table.columns)
                fields = [("_record", "層級"), ("row_label", "列階層"), ("column_label", "欄階層／系列"), (metric, METRICS[metric].replace("（比例）", "（%）")), ("row_count", "紀錄數")]
                widths = [55, 200, 200, width-500, 45]
                if comparison:
                    fields += [(metric + "_previous", "前期值"), (metric + "_delta", "差異（百分點）" if metric == "net_margin_rate" else "差額"), *([] if metric == "net_margin_rate" else [(metric + "_growth", "成長率")]), (metric + "_status", "比較狀態")]
                    widths = [40, 120, 100, 78, 40, 78, 78, *([] if metric == "net_margin_rate" else [70]), 100]
                    factor = width / sum(widths); widths = [value * factor for value in widths]
                story.append(paragraph(METRICS[metric].replace("（比例）", "（%）") + " · 完整彙總", "card"))
                cells = [[paragraph(label, "header") for _, label in fields]]
                commands = [("BACKGROUND", (0,0), (-1,0), PURPLE), ("VALIGN", (0,0), (-1,-1), "TOP"), ("BOTTOMPADDING", (0,0), (-1,-1), 5), ("TOPPADDING", (0,0), (-1,-1), 5), ("LINEBELOW", (0,0), (-1,0), .5, PURPLE), ("LINEBELOW", (0,1), (-1,-1), .25, colors.HexColor("#e5e7ed"))]
                for row_index, record in enumerate(data, 1):
                    row = []
                    for key, _ in fields:
                        value = record.get(key)
                        if key == metric or key == metric + "_previous": text = exact(value, metric, table.card)
                        elif key == metric + "_delta": text = "無比較差額" if value is None else f"{value:+,.2f}" + (" 個百分點" if metric == "net_margin_rate" else "")
                        elif key == metric + "_growth": text = "無法計算" if value is None else f"{value * 100:+,.2f}%"
                        else: text = "" if value is None else f"{value:,}" if isinstance(value, int) else str(value)
                        row.append(paragraph(text, "small"))
                    cells.append(row)
                    if record.get("_record") in ("小計", "總計"): commands.append(("BACKGROUND", (0,row_index), (-1,row_index), PALE))
                grid = LongTable(cells, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=1)
                grid.setStyle(TableStyle(commands))
                story += [grid, Spacer(1, 16)]
    doc = SimpleDocTemplate(output, pagesize=page_size, rightMargin=36, leftMargin=36, topMargin=38, bottomMargin=35, title=body.config.name, author="報表工作區", pageCompression=1)
    def footer(canvas, document):
        canvas.setFont(FONT, 8); canvas.setFillColor(colors.HexColor("#647084"))
        canvas.drawString(36, 19, pdf_text("報表工作區 · 完整報表 · " + body.config.name[:40]))
        canvas.drawRightString(page_size[0]-36, 19, f"第 {document.page} 頁")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output
